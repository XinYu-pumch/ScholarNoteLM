import copy,json,sys,unittest,zipfile
from unittest.mock import patch
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import review_bundle as b,review_content as c,review_export as x,review_format as fmt
import test_review_bundle as fixtures
import test_review_writing as writing_fixture
import review_writing as w,engine as e,evidence_resources,vector_chunks
from test_evidence_resources import PNG

class ResourceBlocksTests(unittest.TestCase):
 def setUp(self):self.fx=fixtures.BundleTests();self.fx.setUp()
 def tearDown(self):self.fx.tearDown()
 def value(self,blocks):
  value=self.fx.output();value['sections'][1]['blocks']=blocks;return value
 def decode(self,blocks,profile=None):return b.decode(self.value(blocks),self.fx.write(),self.fx.unit,profile)
 def test_contiguous_citations_collapse_but_all_ids_and_positions_survive(self):
  text='Claim [E1] [R1] [R2] [E2] [E1]。 Another [E1]。\n\nAgain [E1].'
  sections,errors=self.decode([{'type':'markdown','text':text}]);self.assertFalse(errors);s=sections[1]
  self.assertEqual(s['markdown'],'Claim (PMID: 12345678, 23456789)。 Another (PMID: 12345678)。\n\nAgain (PMID: 12345678).')
  self.assertEqual(s['sourceBlocks'][0]['text'],text)
  self.assertEqual([r['id'] for r in s['citationRefs']],['E1','R1','R2','E2'])
  self.assertEqual(len(s['blocks']),1)  # A resource citation never auto-inserts the resource.
 def test_deduplication_does_not_cross_clauses_cells_or_lines(self):
  self.assertEqual(c.dedupe_citations('(PMID: 1), (PMID: 2) (PMID: 1)'),'('+ 'PMID: 1, 2)')
  for sep in ('。','; ','\n','\n\n',' | ',' and '):
   text='(PMID: 1)'+sep+'(PMID: 1)';self.assertEqual(c.dedupe_citations(text),text)
 def test_adjacent_and_grouped_ids_keep_all_bindings_and_collapse_display(self):
  for text in ('Claim [E1][R1][E2]。 Another [E1]。','Claim [E1, R1, E2]。 Another [E1]。'):
   sections,errors=self.decode([{'type':'markdown','text':text}]);self.assertFalse(errors,errors)
   self.assertEqual([r['id'] for r in sections[1]['citationRefs']],['E1','R1','E2'])
   self.assertEqual(sections[1]['sourceBlocks'][0]['text'],text)
   self.assertEqual(sections[1]['markdown'],'Claim (PMID: 12345678, 23456789)。 Another (PMID: 12345678)。')
 def test_link_syntax_and_unknown_group_members_remain_rejected(self):
  for text in ('Claim [E1](https://example.org).','Claim [E1][link].','[E1]: https://example.org','Claim [E1, E999].','Claim [E1][R999].','Claim [E1, invented].'):
   self.assertTrue(self.decode([{'type':'markdown','text':text}])[1],text)
 def test_only_selected_resource_is_embedded_and_numbering_is_stable(self):
  sections,errors=self.decode([{'type':'markdown','text':'Intro [R2]'}, {'type':'resource','resource_id':'R1'}, {'type':'markdown','text':'After [E1]'}]);self.assertFalse(errors)
  rendered=c.render_sections(sections,self.fx.bundle['resources'],self.fx.project)
  blocks=rendered[1]['displayBlocks'];self.assertEqual([v['type'] for v in blocks],['markdown','resource','markdown'])
  self.assertEqual(blocks[1]['label'],'表 1');self.assertIn('42',blocks[1]['html']);self.assertFalse(any(v.get('dataUrl') for v in blocks))
  self.assertEqual(sections[1]['sourceBlocks'][1],{'type':'resource','resource_id':'R1'})
  self.assertEqual(sections[1]['blocks'][1]['resourceId'],self.fx.objects[0]['id'])
 def test_figure_requires_actual_request_image_availability(self):
  blocks=[{'type':'resource','resource_id':'R2'}]
  for profile in ({},{'model':'text','visionModels':['vision']},{'model':'vision','visionModels':[]}):
   self.assertIn('a',self.decode(blocks,profile)[1])
  package=self.fx.write();package['resources'][1].pop('imageFile')
  self.assertIn('a',b.decode(self.value(blocks),package,self.fx.unit,{'model':'vision','visionModels':['vision']})[1])
  sections,errors=self.decode(blocks,{'model':'vision','visionModels':['vision']});self.assertFalse(errors)
  self.assertEqual(sections[1]['blocks'][0]['imageInput'],'attached')
 def test_unknown_wrong_type_extra_fields_and_intro_resources_rejected(self):
  for rid in ('R999','E1','../outside'):
   self.assertIn('a',self.decode([{'type':'resource','resource_id':rid}])[1])
  with self.assertRaises(ValueError):self.decode([{'type':'resource','resource_id':'R1','path':'/tmp/outside'}])
  value=self.fx.output();value['sections'][0]['blocks']=[{'type':'resource','resource_id':'R1'}]
  self.assertIn('parent',b.decode(value,self.fx.write(),self.fx.unit)[1])
  self.assertIn('a',self.decode([{'type':'markdown','text':'![invented](file:///tmp/x.png)'}])[1])
 def test_synthesized_partial_or_merged_table_is_regular_markdown(self):
  self.fx.bundle['schemaVersion']='review-unit-resource-blocks-v2'
  text='| Source | Selected result |\n|---|---|\n| A | 42 [E1] [R1] |\n| B | Other [E2] |'
  sections,errors=self.decode([{'type':'markdown','text':text}]);self.assertFalse(errors)
  blocks=c.render_sections(sections,self.fx.bundle['resources'],self.fx.project)[1]['displayBlocks']
  self.assertEqual(blocks[0]['type'],'notice');self.assertEqual(blocks[1]['type'],'markdown')
  html=str(x.blocks_html(blocks));self.assertEqual(html.count('<table>'),1);self.assertNotIn('<figure',html)
  self.assertEqual(html.count('PMID: 12345678'),1);self.assertIn('PMID: 23456789',html)
 def test_offline_html_word_and_markdown_keep_selected_assets(self):
  chosen=[{'type':'resource','resource_id':'R1'},{'type':'resource','resource_id':'R2'},{'type':'markdown','text':'Claim [E1] [R1].'}]
  sections,errors=self.decode(chosen,{'model':'vision','visionModels':['vision']});self.assertFalse(errors)
  rendered=c.render_sections(sections,self.fx.bundle['resources'],self.fx.project)
  body={n['serial']:s['markdown'] for n,s in zip(self.fx.nodes,rendered)};blocks={n['serial']:s['displayBlocks'] for n,s in zip(self.fx.nodes,rendered)}
  html=x.html_document('Test',self.fx.nodes,body,{},False,blocks)
  from bs4 import BeautifulSoup
  soup=BeautifulSoup(html,'html.parser');content=soup.select_one('#s-1\\.1 .content')
  self.assertEqual(len(content.select('figure')),2);self.assertEqual(len(content.select('table')),1)
  pool=json.loads(soup.find(id='reference-data').string)
  self.assertTrue(pool['images'][content.img['data-image-ref']].startswith('data:image/png;base64,'));self.assertIn('42',content.get_text())
  doc=self.fx.project/'test.docx';x.word_document(doc,'Test',self.fx.nodes,body,False,blocks)
  with zipfile.ZipFile(doc) as archive:
   self.assertTrue(any(p.startswith('word/media/') for p in archive.namelist()));self.assertIn('42',archive.read('word/document.xml').decode())
  md=c.markdown(blocks['1.1'],lambda block:'resources/figure.png');self.assertIn('![图 1](resources/figure.png)',md);self.assertIn('| A | 42 |',md)
 def test_legacy_saved_public_format_and_evidence_are_preserved(self):
  # Saved sections may lack blocks (markdown-only); their stored format and evidence must survive.
  value={'unit_id':'1','sections':[{'node_id':n['serial'],'title':n['title'],'blocks':[{'type':'markdown','text':'Intro' if i==0 else 'Claim [E1]'}]} for i,n in enumerate(self.fx.nodes)]}
  sections,errors=b.decode(value,self.fx.bundle,self.fx.unit);self.assertFalse(errors)
  for s in sections:s.pop('blocks');s.pop('sourceBlocks')
  sections[1]['markdown']='Claim (PMID: 12345678) (PMID: 12345678)'
  public=fmt.public_part(self.fx.unit,sections,self.fx.evidence);before=copy.deepcopy(public)
  rendered=c.render_sections(sections,self.fx.bundle['resources'],self.fx.project)
  self.assertEqual(rendered[1]['markdown'],'Claim (PMID: 12345678)');self.assertEqual(public,before)
  self.assertEqual(public['format'],self.fx.bundle['schemaVersion'])
 def test_same_local_id_across_units_does_not_alias_resource(self):
  sections,_=self.decode([{'type':'resource','resource_id':'R1'}]);numbers={}
  first=c.render_sections(sections,self.fx.bundle['resources'],self.fx.project,numbers)
  other=copy.deepcopy(self.fx.bundle['resources']);other[0]['resourceId']='another';changed=copy.deepcopy(sections);changed[1]['blocks'][0]['resourceId']='another'
  second=c.render_sections(changed,other,self.fx.project,numbers)
  self.assertEqual(first[1]['displayBlocks'][0]['label'],'表 1');self.assertEqual(second[1]['displayBlocks'][0]['label'],'表 2')
  again=c.render_sections(sections,self.fx.bundle['resources'],self.fx.project,numbers);self.assertEqual(again[1]['displayBlocks'][0]['label'],'表 1')
 def test_original_html_table_spans_survive_word_export(self):
  resource={'type':'resource','resourceId':'table','label':'表 1','kind':'table','pmid':'12345678','content':'<table><tr><th rowspan="2">Group</th><th colspan="2">Results</th></tr><tr><td>A</td><td>B</td></tr><tr><td>X</td><td>1</td><td>2</td></tr></table>'}
  dest=self.fx.project/'spans.docx';x.word_document(dest,'Test',self.fx.nodes,{'1.1':'Table (PMID: 12345678)'},False,{'1.1':[resource]})
  from docx import Document
  table=Document(dest).tables[0]
  self.assertEqual(table.cell(1,0).text,'Group');self.assertEqual(table.cell(0,2).text,'Results');self.assertEqual(table.cell(1,2).text,'B');self.assertEqual(table.cell(2,2).text,'2')

class ResourceWorkerTests(unittest.TestCase):
 setUp=writing_fixture.WritingTests.setUp
 tearDown=writing_fixture.WritingTests.tearDown
 setup_run=writing_fixture.WritingTests.setup_run
 run_both=writing_fixture.WritingTests.run_both
 status=writing_fixture.WritingTests.status
 def test_worker_keeps_blocks_and_all_exports_including_external_md_images(self):
  e.atomic(self.base/'markdown/assets/12345678/figure.png',PNG)
  text='Before.\n\n| Item | Value |\n|---|---|\n| A | 42 |\n\n![](assets/12345678/figure.png)\n\nAfter.'
  objects=evidence_resources.extract(self.base,'12345678','v',text,vector_chunks.split(text),e.atomic)
  record={'id':'one','text':'Claim','metadata':{'pmid':'12345678','title':'Study','sourceKind':'fulltext','sectionPath':[],'documentVersion':'v'},'resources':objects}
  class Retriever:
   def __init__(self,*args):pass
   def retrieve(self,*args):return [record]
  async def chat(profile,messages,cancelled,progress=lambda **kw:None):
   content=messages[-1]['content'];self.assertEqual(sum(p['type']=='image_url' for p in content),1)
   data=json.loads(content[0]['text']);value=data['outputContract']
   for section in value['sections']:
    role=next(n['role'] for n in data['unit']['nodes'] if n['node_id']==section['node_id'])
    section['blocks']=[{'type':'markdown','text':'Introduction'}] if role=='introduction' else [{'type':'markdown','text':'Claim [E1] [R1].'},{'type':'resource','resource_id':'R1'},{'type':'resource','resource_id':'R2'}]
   return json.dumps(value)
  with patch.dict(writing_fixture.PROFILE,{'visionModels':['fixture']}),patch.object(w.retrieval,'corpus',return_value=[record]),patch.object(w,'chat',side_effect=chat):
   self.run_both(self.setup_run(),Retriever)
   self.assertEqual(self.status()['latest'],4)
   preview=w.rpc(e,'Test','preview',{})
   self.assertEqual(sum(b['type']=='resource' for b in preview['displayBlocks']),10)
   self.assertEqual({b['label'] for b in preview['displayBlocks'] if b['type']=='resource'},{'表 1','图 1'})
   for form in ('md','html','docx'):w.export(e,'Test',form)
   html=next((self.path/'exports').glob('*/*.html')).read_text();self.assertIn('class="body-resource"',html);self.assertIn('data:image/png;base64,',html)
   md=next((self.path/'exports').glob('*/*.md')).read_text();self.assertIn('![图 1](resources/',md)
   self.assertEqual(len(list((self.path/'exports').glob('*/resources/*.png'))),1)
   self.assertNotIn('(PMID: 12345678) (PMID: 12345678)',md)
   saved=json.loads(next((self.path/'review_parts').glob('section_*.json')).read_text());self.assertIn('sourceBlocks',saved['sections'][0]);self.assertEqual(saved['format'],b.VERSION)
   archive=w.manuscripts.save(e,'Test');w.rpc(e,'Test','reset',{'confirm':True})
   import shutil
   shutil.rmtree(self.base/'markdown')
   for form in ('md','html','docx'):
    result=w.rpc(e,'Test','manuscript.export',{'id':archive['id'],'format':form});self.assertTrue(Path(result['path']).is_file())
    if form=='html':self.assertIn('data:image/png;base64,',Path(result['path']).read_text())

if __name__=='__main__':unittest.main()
