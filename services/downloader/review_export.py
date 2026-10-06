"""Standalone HTML reader and Word, fed by the same normalized sections."""
import html,json,re,time,sys
from pathlib import Path
from urllib.parse import urlparse
import review_format as fmt
import review_content
import html_compaction
from text_presentation import strip_source_links

def markdown_html(text):
    import markdown
    from bs4 import BeautifulSoup
    soup=BeautifulSoup(markdown.markdown(html.escape(strip_source_links(review_content.dedupe_citations(text))),extensions=['tables','fenced_code','sane_lists']),'html.parser')
    # Source markup is data. Keep Markdown formatting without active HTML or URLs.
    for a in soup.find_all('a'):
        url=a.get('href','');scheme=urlparse(url).scheme.lower()
        if url.lower().startswith('#page-'):
            a.unwrap();continue
        if scheme not in ('http','https','mailto',''):a.attrs.pop('href',None)
        a['rel']='noopener noreferrer'
    for img in soup.find_all('img'):img.replace_with(img.get('alt',''))
    return soup

def resource_html(text):
    """Render converted HTML tables without copying active tags or attributes."""
    if '<table' not in text.lower():return str(markdown_html(text))
    from bs4 import BeautifulSoup
    source=BeautifulSoup(text,'html.parser');parts=[]
    for table in source.find_all('table'):
        if table.find_parent('table'):continue
        rows=[]
        for row in table.find_all('tr'):
            cells=[]
            for cell in row.find_all(['td','th'],recursive=False):
                spans=''
                for key in ('rowspan','colspan'):
                    value=str(cell.get(key,'1'))
                    if value.isdigit():spans+=f' {key}="{max(1,min(1000,int(value)))}"'
                cells.append('<'+cell.name+spans+'>'+html.escape(cell.get_text('\n'))+'</'+cell.name+'>')
            rows.append('<tr>'+''.join(cells)+'</tr>')
        parts.append('<table><tbody>'+''.join(rows)+'</tbody></table>')
    return ''.join(parts)

def reference_entry(chunk,project,seen):
    import evidence_resources
    objects=[]
    for obj in chunk.get('resources',[]):
        if obj['id'] in seen:continue
        seen.add(obj['id']);entry={'id':obj['id'],'kind':obj['kind'],'pmid':obj['pmid'],
            'lineStart':obj['lineStart'],'lineEnd':obj['lineEnd'],
            'html':resource_html(obj['content']) if obj['kind']=='table' else '',
            'context':[{'position':'before' if p['charEnd']<=obj['charStart'] else 'after','html':str(markdown_html(p['content']))} for p in obj.get('context',[])]}
        if obj['kind']=='figure':
            if obj.get('assetPath'):entry['dataUrl']=evidence_resources.image_data(project,obj)
            else:entry['error']=obj.get('imageError','图片快照不可用')
        objects.append(entry)
    return {'id':chunk['id'],'html':str(markdown_html(chunk['text'])),'resources':objects}

def generated_table_html(block):
    def inline(text):
        soup=markdown_html(text)
        for p in soup.find_all('p'):p.unwrap()
        return str(soup).replace('\n','<br>')
    header=''.join('<th>'+inline(c)+'</th>' for c in block['columns'])
    rows=''.join('<tr>'+''.join('<td>'+inline(c)+'</td>' for c in row)+'</tr>' for row in block['rows'])
    return '<table><thead><tr>'+header+'</tr></thead><tbody>'+rows+'</tbody></table>'

def blocks_html(blocks):
    from bs4 import BeautifulSoup
    parts=[]
    for block in blocks:
        if block['type']=='markdown':parts.append(str(markdown_html(block['text'])));continue
        if block['type']=='notice':parts.append('<aside class="object-notice">'+html.escape(block['text'])+'</aside>');continue
        anchor=' id="'+html.escape(block['anchor'],quote=True)+'"' if block.get('anchor') else ''
        if block['type']=='table':
            caption=str(markdown_html(block['label']+' · '+block['title']))
            notes=str(markdown_html(block['notes'])) if block['notes'] else ''
            parts.append('<figure class="body-table"'+anchor+'><figcaption>'+caption+'</figcaption>'+generated_table_html(block)+'<div class="table-notes">'+notes+'</div></figure>');continue
        label=html.escape(block['label']);pmid=html.escape(block['pmid'])
        if block['kind']=='table':content=resource_html(block['content'])
        else:
            data=block.get('dataUrl','')
            if not re.fullmatch(r'data:image/(?:png|jpeg|gif|webp);base64,[A-Za-z0-9+/=]+',data):raise ValueError('正文图片快照格式无效')
            content=f'<img src="{data}" alt="{label} · PMID {pmid}" loading="lazy">'
        context=''.join(str(markdown_html(p['text'])) for p in block.get('context',[]))
        details='<details class="body-resource-context"><summary>原文邻近文字（不保证为完整图注）</summary>'+context+'</details>' if context else ''
        parts.append(f'<figure class="body-resource"{anchor} data-resource-id="{html.escape(block["resourceId"],quote=True)}">{content}<figcaption>{label} · 原文资源 (PMID: {pmid})</figcaption>{details}</figure>')
    return BeautifulSoup(''.join(parts),'html.parser')

def html_document(topic,rows,sections,references,partial,section_blocks=None):
    from bs4 import NavigableString
    toc=[];body=[]
    for index,row in enumerate(rows):
        sid=row['serial'];safe_sid=html.escape(sid,quote=True);title=html.escape(row['title']);depth=len(sid.split('.'));level=min(depth,6)
        has_children=index+1<len(rows) and rows[index+1]['serial'].startswith(sid+'.')
        toggle=(f'<button class="toc-toggle" type="button" aria-expanded="true" aria-label="折叠 {safe_sid}">⌄</button>' if has_children else '<span class="toc-spacer"></span>')
        toc.append(f'<div class="toc-item" data-serial="{safe_sid}" data-depth="{depth}" style="--indent:{min(depth-1,5)*12}px">{toggle}<a href="#s-{safe_sid}"><span class="toc-number">{safe_sid}</span><span>{title}</span></a></div>')
        body.append(f'<section id="s-{safe_sid}" class="sec depth-{level}"><h{level}><span class="section-number">{safe_sid}</span> {title}</h{level}>')
        text=sections.get(sid,'')
        if text:
            soup=blocks_html(section_blocks[sid]) if section_blocks is not None and sid in section_blocks else markdown_html(text)
            for node in list(soup.find_all(string=True)):
                if node.parent.name in ('code','pre','a'):continue
                matches=list(re.finditer(r'[（(]\s*PMID\s*[:：]\s*([^）)]+)[）)]',str(node),re.I))
                if not matches:continue
                cursor=0
                for match in matches:
                    node.insert_before(NavigableString(str(node)[cursor:match.start()]));node.insert_before(NavigableString('(PMID: '))
                    for i,pmid in enumerate(re.findall(r'\d+',match[1])):
                        if i:node.insert_before(NavigableString(', '))
                        button=soup.new_tag('button',attrs={'class':'lnk','type':'button','data-pmid':pmid,'data-scope':sid});button.string=pmid;node.insert_before(button)
                    node.insert_before(NavigableString(')'));cursor=match.end()
                node.insert_before(NavigableString(str(node)[cursor:]));node.extract()
            body.append('<div class="content">'+str(soup)+'</div>')
        body.append('</section>')
    data=json.dumps(references,ensure_ascii=False).replace('<','\\u003c').replace('>','\\u003e').replace('&','\\u0026')
    css=Path(__file__).with_name('review_export.css').read_text(encoding='utf8')
    script=Path(__file__).with_name('review_export.js').read_text(encoding='utf8')
    pmids=set(re.findall(r'data-pmid="(\d+)"',''.join(body)))
    stamp=time.strftime('%Y-%m-%d %H:%M')
    document = f'''<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>{html.escape(topic)}</title><style>{css}</style></head>
<body class="toc-open">
<a class="skip-link" href="#main">跳至正文</a>
<header class="reader-bar"><div class="brand"><span class="brand-mark" aria-hidden="true">S</span>ScholarNoteLM<span class="brand-divider">/</span><span class="bar-title">综述阅读</span></div>
<div class="reader-actions"><button id="toggle-toc" type="button" aria-controls="toc" aria-expanded="true">目录</button><button id="toggle-ref" type="button" aria-controls="ref" aria-expanded="false">引用详情</button><button id="print" type="button">打印 / PDF</button></div></header>
<div class="reading-progress" aria-hidden="true"><span id="progress"></span></div>
<button id="scrim" type="button" aria-label="关闭侧栏" hidden></button>
<nav id="toc" aria-label="目录"><div class="panel-heading"><div><span class="eyebrow">CONTENTS</span><h2>目录 <span class="count">{len(rows)}</span></h2></div><button class="panel-close" data-close="toc" type="button" aria-label="关闭目录">×</button></div><div class="toc-list">{''.join(toc)}</div><div class="toc-footer">阅读进度 <span id="progress-label">0%</span></div></nav>
<main id="main" tabindex="-1"><article class="document"><header class="document-header"><div class="eyebrow">LITERATURE REVIEW <span>文献综述</span></div><h1>{html.escape(topic)}</h1><div class="document-meta"><span>{len(rows)} 个章节</span><span>{len(pmids)} 篇引用文献</span><time>导出于 {stamp}</time></div>{'<p class="draft">草稿 · 含未完成或待更新章节</p>' if partial else ''}</header>
{''.join(body)}<footer class="document-footer"><span>ScholarNoteLM</span><span>文献综述 · {stamp}</span></footer></article></main>
<aside id="ref" aria-label="引用详情" hidden><div class="panel-heading"><div><span class="eyebrow">SOURCE MATERIAL</span><h2>引用详情</h2></div><button class="panel-close" data-close="ref" type="button" aria-label="关闭引用详情">×</button></div><div id="area" aria-live="polite"><div class="tip"><span class="tip-symbol" aria-hidden="true">↗</span><strong>从引用回到原文</strong><p>点击正文中的 PMID，查看该章节使用的文献素材。</p></div></div></aside>
<script id="reference-data" type="application/json">{data}</script><script>{script}</script></body></html>'''
    return html_compaction.compact_html(document)

def word_document(dest,topic,rows,sections,partial,section_blocks=None):
    from docx import Document
    from docx.shared import Pt,Inches
    import base64,io
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from bs4 import NavigableString
    doc=Document();normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(11)
    cjk='Songti SC' if sys.platform=='darwin' else 'Microsoft YaHei'
    for style_name in ['Normal','Title',*[f'Heading {i}' for i in range(1,10)]]:
        style=doc.styles[style_name];style.font.name='Arial';style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),cjk)
    title=doc.add_heading(topic,0);title.alignment=WD_PARAGRAPH_ALIGNMENT.CENTER
    doc.add_paragraph('生成时间：'+time.strftime('%Y-%m-%d %H:%M'))
    if partial:doc.add_paragraph('草稿：含未完成或待更新章节。')
    doc.add_page_break()
    def inline(p,node,bold=False,italic=False):
        if isinstance(node,NavigableString):
            run=p.add_run(str(node));run.bold=bold;run.italic=italic;return
        if node.name=='br':p.add_run().add_break();return
        if node.name=='a' and re.fullmatch(r'#obj_[a-f0-9]{24}',node.get('href','')):
            link=OxmlElement('w:hyperlink');link.set(qn('w:anchor'),node['href'][1:]);run=OxmlElement('w:r');value=OxmlElement('w:t');value.text=node.get_text();run.append(value);link.append(run);p._p.append(link);return
        for child in node.children:inline(p,child,bold or node.name in ('strong','b'),italic or node.name in ('em','i'))
    bookmarks=set()
    def blocks(soup):
        for node in soup.children:
            if isinstance(node,NavigableString):
                if str(node).strip():doc.add_paragraph(str(node))
            elif node.name in ('ul','ol'):
                for li in node.find_all('li',recursive=False):inline(doc.add_paragraph(style='List Bullet' if node.name=='ul' else 'List Number'),li)
            elif node.name=='table':
                trs=node.find_all('tr');occupied=set();placements=[];width=0
                for row,tr in enumerate(trs):
                    col=0
                    for td in tr.find_all(['th','td'],recursive=False):
                        span=max(1,int(td.get('colspan',1)));height=min(len(trs)-row,max(1,int(td.get('rowspan',1))))
                        while any((row,c) in occupied for c in range(col,col+span)):col+=1
                        placements.append((row,col,height,span,td))
                        occupied.update((r,c) for r in range(row,row+height) for c in range(col,col+span));col+=span;width=max(width,col)
                if width:
                    table=doc.add_table(rows=len(trs),cols=width);table.style='Table Grid'
                    for row,col,height,span,td in placements:
                        cell=table.cell(row,col)
                        if height>1 or span>1:cell=cell.merge(table.cell(row+height-1,col+span-1))
                        inline(cell.paragraphs[0],td,td.name=='th')
            elif node.name=='figure':
                ident=node.get('id','')
                if re.fullmatch(r'obj_[a-f0-9]{24}',ident) and ident not in bookmarks:
                    p=doc.add_paragraph();start=OxmlElement('w:bookmarkStart');end=OxmlElement('w:bookmarkEnd');index=str(len(bookmarks)+1)
                    start.set(qn('w:id'),index);start.set(qn('w:name'),ident);end.set(qn('w:id'),index);p._p.extend([start,end]);bookmarks.add(ident)
                blocks(node)
            elif node.name=='img':
                raw=base64.b64decode(node['src'].split(';base64,',1)[1])
                if node['src'].startswith('data:image/webp;'):
                    from PIL import Image
                    converted=io.BytesIO();Image.open(io.BytesIO(raw)).save(converted,format='PNG');raw=converted.getvalue()
                doc.add_picture(io.BytesIO(raw),width=Inches(6))
            elif node.name=='details':continue
            else:inline(doc.add_paragraph(),node)
    used=set()
    for row in rows:
        sid=row['serial'];doc.add_heading(sid+' '+row['title'],min(len(sid.split('.')),9))
        text=sections.get(sid,'')
        if text:blocks(blocks_html(section_blocks[sid]) if section_blocks is not None and sid in section_blocks else markdown_html(text));used.update(re.findall(r'\d+',','.join(re.findall(r'PMID\s*[:：]\s*([\d\s,，、;；]+)',text,re.I))))
    if used:
        doc.add_page_break();doc.add_heading('参考文献',1)
        for pmid in sorted(used,key=int):
            p=doc.add_paragraph();p.add_run(f'[PMID: {pmid}] ').bold=True
            link=OxmlElement('w:hyperlink');link.set(qn('r:id'),doc.part.relate_to(f'https://pubmed.ncbi.nlm.nih.gov/{pmid}/','http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True))
            run=OxmlElement('w:r');text=OxmlElement('w:t');text.text=f'https://pubmed.ncbi.nlm.nih.gov/{pmid}/';run.append(text);link.append(run);p._p.append(link)
    temp=dest.with_suffix('.tmp');doc.save(temp);temp.replace(dest)
